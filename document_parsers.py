# document_parsers.py
from PyPDF2 import PdfReader
import docx
import openpyxl
import pptx

def parse_pdf(file):
    """
    Parses a PDF file and returns its text content.
    """
    reader = PdfReader(file)
    text = ""
    for page in reader.pages:
        text += page.extract_text()
    return text

def parse_docx(file):
    """
    Parses a DOCX file and returns its text content.
    """
    doc = docx.Document(file)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text

def parse_xlsx(file):
    """
    Parses an XLSX file and returns its text content.
    """
    workbook = openpyxl.load_workbook(file)
    text = ""
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                text += str(cell.value) + " "
            text += "\n"
    return text

def parse_pptx(file):
    """
    Parses a PPTX file and returns its text content.
    """
    prs = pptx.Presentation(file)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text
